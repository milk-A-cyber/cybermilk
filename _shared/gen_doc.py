# -*- coding: utf-8 -*-
"""Generate 可视化数据分析网站-关键亮点说明.docx (DataFrame Studio)."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

OUT = r"E:\qq\s\可视化导论\3-数据分析\index\DataFrame-Studio-可视化数据分析网站-关键亮点说明.docx"

BLACK = RGBColor(0, 0, 0)


def set_run(run, latin="Times New Roman", east="宋体", size=12, bold=False, black=False):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.font.bold = bold
    if black:
        run.font.color.rgb = BLACK


def style_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0,
               line=1.5, indent_pt=None):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    else:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if indent_pt is not None:
        pf.first_line_indent = Pt(indent_pt)


def add_title(doc, text):
    p = doc.add_paragraph(style="Title")
    style_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=18, line=1.0)
    set_run(p.add_run(text), latin="Times New Roman", east="黑体", size=18, bold=True, black=True)


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    style_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=12, line=1.0)
    set_run(p.add_run(text), latin="Times New Roman", east="楷体", size=12, black=True)


def add_h1(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    style_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=14, after=6, line=1.0)
    set_run(p.add_run(text), latin="Times New Roman", east="黑体", size=16, bold=True, black=True)


def add_h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    style_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=11, after=5, line=1.0)
    set_run(p.add_run(text), latin="Times New Roman", east="黑体", size=14, bold=True, black=True)


def add_body(doc, text):
    p = doc.add_paragraph()
    style_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, line=1.5, indent_pt=24)
    set_run(p.add_run(text), latin="Times New Roman", east="宋体", size=12, black=True)
    return p


doc = Document()

# --- page setup: A4, margins 2.5 cm ---
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.5)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)

# --- default Normal style ---
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
nrpr = normal.element.get_or_add_rPr()
nrfonts = nrpr.get_or_add_rFonts()
nrfonts.set(qn("w:ascii"), "Times New Roman")
nrfonts.set(qn("w:hAnsi"), "Times New Roman")
nrfonts.set(qn("w:eastAsia"), "宋体")

add_title(doc, "DataFrame Studio · 可视化 pandas 数据分析工作台")
add_subtitle(doc, "—— 以可视化图形增强数据分析能力的双语交互网站关键亮点说明")

# 一、网站概述
add_h1(doc, "一、网站概述")
add_body(doc, "本网站是“可视化·数据分析”课程作业成果，是一个以 Python 的 pandas 库为参照、以可视化为表达方式的数据分析工作台。全站内容中英双语，右上角一键切换，两种语言在内容、功能与数据上完全一致。")
add_body(doc, "网站可自由加载数据：支持拖拽或选择本地的 CSV、TSV、TXT、JSON 文件直接载入；同时内置三套示例数据，并默认加载“季度销售订单”（340 条订单流水，含地区、品类、渠道、折扣与利润），打开即用、即选即析；另备鸢尾花 Iris（150 行经典生物测量数据）与乘客名单（240 条、年龄与票价刻意缺失）两套数据，分别用于演示分组透视与缺失值处理。")
add_body(doc, "网站以 pandas 为核心参照，把数据读取、选择列、筛选、排序、缺失值处理、去重、计算列、列操作、透视表、描述统计、频数统计、相关性与图表绘制等功能逐一“图形化”实现，并以图表把每一步分析结论直观呈现，让抽象的数据操作变得“所见即所得”。")

# 二、整体结构与版式
add_h1(doc, "二、整体结构与版式")
add_h2(doc, "（一）四段式功能分组")
add_body(doc, "网站按 pandas 的分析流程把全部功能组织为四个分组：数据（数据源、数据表），负责加载、切换与浏览；转换（选择列、筛选、排序、缺失值、去重、计算列、列操作、透视表），负责清洗与变换；分析（概览、描述统计、频数统计、相关性、图表工坊），负责统计与可视化；输入/输出（导出），负责结果的保存与代码交付。")
add_h2(doc, "（二）操作流水线")
add_body(doc, "网站把“每次转换”串成一条可视化的操作流水线：任何一次选择列、筛选、排序、去重或计算列都会以一块“流水线片”加入顶栏并逐级向后传递，等效于 pandas 的链式操作。用户可以随时点击移除某一步或一键清空流水线，底部状态栏同步显示当前数据的行列数与已执行步骤数，让每一步对数据的改变都可回溯、可理解。")
add_h2(doc, "（三）版式设计")
add_body(doc, "左侧为分组导航栏，内容区按卡片分区，各功能相互独立、浏览路径清晰；顶部为工具条，含当前数据集徽标、明/暗主题切换、一键重置与中英语言切换；底部为状态栏。整站支持明暗双主题与响应式布局，配色克制、留白充分，保证“美观耐看”。")
add_h2(doc, "（四）三大设计目标")
add_body(doc, "功能强大完整：覆盖 pandas 从数据加载到清洗、变换、统计、可视化再到导出的完整链路；交互方便简洁：全程点选、拖动、滑杆即可完成，无需书写代码；页面美观耐看：图表风格统一、卡片化布局、双主题适配。")

# 三、各模块的可视化亮点介绍
add_h1(doc, "三、各模块的可视化亮点介绍")
add_h2(doc, "（一）数据源与数据表")
add_body(doc, "数据源提供拖拽上传区与三张内置数据集卡片，并允许设定默认数据；载入后即时显示行列规模与字段概览。数据表将数据以表格呈现，为数值与文本字段打上类型徽标，支持滚动浏览大量行，点击列头即可查看该列的统计与分布。")
add_h2(doc, "（二）概览")
add_body(doc, "概览把“一眼看懂数据”做到极致：以横向条形图展示各列数据类型（dtype）分布，以柱状图呈现各列缺失值数量，并为每个数值列附上迷你直方图，配合“非空值、唯一值、均值、标准差、极值”等统计卡，让数据的整体形态无需读表即了然于胸。")
add_h2(doc, "（三）描述统计")
add_body(doc, "描述统计对应 pandas 的 describe()：除给出 count、mean、std、四分位、median、max 等指标表外，更以“直方图＋核密度、箱线图、累积分布曲线（CDF）”三图联动呈现所选列的分布。鼠标在直方图上悬停或框选，即可聚焦对应数值区间并实时显示该区间占比；箱线图同步标出中位数、四分位距（IQR）与离群值，把“统计数字”变成“看得见的分布”。")
add_h2(doc, "（四）频数统计")
add_body(doc, "频数统计对应 value_counts()，将分类字段的取值频数以条形图呈现，支持 TopN 截断与升降序切换，一键识别主导类别及其占比。")
add_h2(doc, "（五）相关性")
add_body(doc, "相关性对应 corr()：以相关矩阵热力图展示数值列两两相关系数，颜色深浅与数值大小一一对应，并支持 pearson、spearman 两种方法切换；下方联动成对散点图，可自由选择 X 轴、Y 轴并按某分类字段分组着色，直观验证“相关”的图形形态与方向。")
add_h2(doc, "（六）图表工坊")
add_body(doc, "图表工坊是“用图形做分析”的核心：支持柱状图、折线图、面积图、饼图、箱线图、直方图六种图形，可指定 X/Y 字段、聚合方式（求和、均值、中位数、计数、最大值、最小值、唯一值）、按分类分组堆叠、TopN 条数（滑杆）与排序方式，图形随配置即时刷新；同时自动生成等价的 pandas 代码（如 groupby(...).mean().plot(kind=“bar”)），让“看得见的图”与“可复现的代码”一一对应。")
add_h2(doc, "（七）缺失值")
add_body(doc, "缺失值模块先以柱状图标出各列缺失数量与占比，让缺口一目了然；提供按前后值填充、按均值/中位数填充等策略，填充前后的缺失分布图同步对比，帮助判断清洗效果。")
add_h2(doc, "（八）透视表")
add_body(doc, "透视表对应 pivot_table()：通过指定行、列、值与聚合方式（均值、求和、计数、中位数、最小值、最大值），即时生成透视结果表，并配以热力图呈现交叉格数值的高低，把“分组汇总”变得直观可读。")
add_h2(doc, "（九）转换类功能")
add_body(doc, "筛选、排序、选择列、计算列、去重、列操作等转换功能均以“可视化条件构造＋实时预览”方式实现：筛选可通过逻辑条件（与/或）与数值区间构造并即时看到命中行数；计算列以函数模板生成新列；选择列、列重命名与删除、去重等操作点选即可完成，每一步结果都会即时反映在数据表与流水线中。")
add_h2(doc, "（十）导出")
add_body(doc, "导出模块将当前（经流水线处理后的）数据一键导出为 CSV 或 JSON，或复制到剪贴板；同时可一键复制整套 pandas 代码，把整条操作流水线“翻译”成可离线运行的 Python 脚本，实现从交互分析到代码交付的无缝衔接。")

# 四、结语
add_h1(doc, "四、结语")
add_body(doc, "网站以“可视化”贯穿数据分析的全过程——用图表替代公式、用点选替代代码、用流水线替代记忆，力求让学习者与使用者“看得懂、点得动、记得住”，直观体会 pandas 数据分析的完整能力。")

doc.save(OUT)
print("SAVED:", OUT, "exists=", os.path.exists(OUT))
